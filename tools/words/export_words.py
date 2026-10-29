# tools/words/export_words.py
# Phase H (#61): collects every line a player reads into one spreadsheet, so Marz can rewrite
# names, flavour, quips and screen text, then hand it back for Claude to apply.
#
#   python tools/words/export_words.py            -> docs/words/pythio-words.xlsx
#
# Needs Python 3 and openpyxl (pip install openpyxl). Run from anywhere; paths are repo-relative.
# Each row keeps the file and line it came from, so the edits can be put back exactly.
# Card rules text ("Heals 32.") is built from the numbers; its wording is listed under
# "Card rules text" with its %s / {TAG} markers, which must stay in any rewrite.
import os
import re
import sys

ROOT = os.path.abspath(os.path.join(os.path.dirname(__file__), "..", ".."))
OUT = os.path.join(ROOT, "docs", "words", "pythio-words.xlsx")

# Areas in the order they appear in the sheet.
AREAS = ["Spells", "Passives", "Battlefield rules", "Pythio's lines", "Fighters and bosses",
         "Title and run screens", "Pick screens", "Practice round", "Escape menu", "HUD",
         "Fighters' warnings", "Tags", "Leaderboard", "Card rules text"]

# Code files with player-facing text, and the area they belong to. The debug panel (F1),
# tests and tools are left out on purpose.
CODE_FILES = {
    "scenes/game.gd": "Title and run screens",
    "scenes/Game.tscn": "Title and run screens",
    "scenes/practice.gd": "Practice round",
    "ui/CardChoicePanel.gd": "Pick screens",
    "ui/PickCards.gd": "Pick screens",
    "ui/PartyChoicePanel.gd": "Pick screens",
    "ui/PartyChoicePanel.tscn": "Pick screens",
    "ui/PauseMenu.gd": "Escape menu",
    "ui/HudCounters.gd": "HUD",
    "ui/LeaderboardView.gd": "Leaderboard",
    "entities/Target.gd": "Fighters' warnings",
    "entities/Target.tscn": "Fighters' warnings",
    "components/statuses/DebuffStatus.gd": "Fighters' warnings",
    "autoload/systems/ContractGenerator.gd": "Fighters and bosses",
    "autoload/systems/BossDirector.gd": "Fighters and bosses",
    "resources/spells/SpellTags.gd": "Tags",
    "resources/spells/SpellData.gd": "Card rules text",
    "resources/spells/passives/TagTriggerRule.gd": "Card rules text",
    "resources/spells/passives/DiminishingShieldsRule.gd": "Battlefield rules",
    "resources/spells/passives/ExhaustionRule.gd": "Battlefield rules",
    "resources/spells/passives/SaturationRule.gd": "Battlefield rules",
}
EFFECTS_DIR = "resources/spells/effects"

# What Pythio's moments mean, for the "Where it shows" column.
MOMENTS = {
    "round_start": "a round starts", "boss": "the Baron shows up", "low_mana": "your mana runs low",
    "death": "a fighter dies", "victory": "a run is won", "spec": "the first spell pick",
    "pick": "a spell pick", "bounty": "the Baron's bounty (passive pick)", "combo": "a combo fires",
    "practice_start": "the practice round starts", "practice_done": "the practice round ends",
    "cleared": "a round is cleared", "skip": "you Skip a pick", "build_complete": "your spell bar is full",
    "debuff_POISON": "Poison first shows up", "debuff_BLEED": "Bleed first shows up",
    "debuff_HEX": "Hex first shows up", "debuff_DREAD": "Dread first shows up",
}

STRING = re.compile(r'"((?:[^"\\]|\\.)*)"')


def rel(path):
    return os.path.relpath(path, ROOT).replace("\\", "/")


def read_lines(path):
    with open(os.path.join(ROOT, path), encoding="utf-8") as f:
        return f.read().split("\n")


def unescape(s):
    return s.replace('\\"', '"').replace("\\n", "\n").replace("\\t", "\t").replace("\\\\", "\\")


def words_in(s):
    """The readable words left once format markers, BBCode and {TAGS} are removed."""
    s = re.sub(r"\[/?[a-z_]+(=[^\]]*)?\]", " ", s)
    s = re.sub(r"%[-+0-9.]*[sdf%]", " ", s)
    s = re.sub(r"\{[A-Z_]+\}", " ", s)
    return re.findall(r"[A-Za-z][A-Za-z']+", s)


def is_player_text(s, line):
    """Heuristic for string literals in code: keep sentences and labels, drop ids and paths."""
    if "://" in s or s.endswith((".png", ".ogg", ".tres", ".gd", ".cfg", ".tscn", ".woff2")):
        return False
    if not words_in(s):
        return False
    if " " not in s and ("/" in s or re.fullmatch(r"[a-z_]+:[a-z_]+", s)):   # paths, "modulate:a"
        return False
    if re.search(r'(==|!=)\s*"' + re.escape(s) + r'"|begins_with\("' + re.escape(s), line):
        return False
    if re.match(r'\s*@export var \w+\s*:\s*String\s*=\s*"', line):     # editor defaults
        return False
    if re.fullmatch(r"[a-z0-9_]+", s):                      # ids, keys, node names
        return False
    if re.fullmatch(r"[A-Za-z0-9_]+", s) and re.search(r"[a-z][A-Z]", s):   # CamelCase node names
        return False
    if re.fullmatch(r"[A-Z0-9_:]+", s):                      # POISON, MENDING, DEBUFF:HEX
        return bool(re.search(r"text|label|title", line, re.I)) and len(s) > 2
    if re.search(r"\b(name|kind|debuff|school|event|action)\s*==?\s*\"" + re.escape(s), line):
        return False
    if re.search(r"(push_warning|push_error|print|assert|get_node|has_status|is_action|"
                 r"_override|get_value|set_value|has_method|call_group|\$)\s*\(?[^\"]*\"" + re.escape(s), line):
        return False
    return True


rows = []   # dicts: area, where, text, file, line


def add(area, where, text, path, line_no):
    rows.append({"area": area, "where": where, "text": text, "file": path, "line": line_no})


def scan_spells():
    folder = os.path.join(ROOT, "resources", "spells")
    for fn in sorted(os.listdir(folder)):
        if not fn.endswith(".tres"):
            continue
        path = "resources/spells/" + fn
        lines = read_lines(path)
        body = "\n".join(lines)
        passive = re.search(r"^passive = ", body, re.M) is not None
        utility = re.search(r"^utility = true", body, re.M) is not None
        area = "Passives" if passive else "Spells"
        name = re.search(r'^display_name = "(.*)"', body, re.M)
        name_line = next((i for i, l in enumerate(lines, 1) if l.startswith("display_name = ")), 1)
        name = name.group(1) if name else fn
        kind = "passive" if passive else ("F-key spell" if utility else "spell")
        for i, l in enumerate(lines, 1):
            m = re.match(r'^(display_name|flavor|combo_label|text_override) = "((?:[^"\\]|\\.)*)"$', l)
            if not m or not m.group(2):
                continue
            key, val = m.group(1), unescape(m.group(2))
            where = {
                "display_name": "Name of the %s (cards, spell bar, tray, tooltips)" % kind,
                "flavor": "%s: flavour line on its card" % name,
                "combo_label": "%s: gold word over the fighter when it fires" % name,
                "text_override": "%s: card rules text, written by hand (keep the {TAG} markers)" % name,
            }[key]
            add(area, where, val, path, i)
        if not re.search(r'^flavor = "..*"$', body, re.M):
            add(area, "%s: no flavour line yet (write one if you like)" % name, "", path, name_line)


def scan_rules():
    folder = os.path.join(ROOT, "resources", "rules")
    for fn in sorted(os.listdir(folder)):
        if fn.endswith(".tres"):
            path = "resources/rules/" + fn
            for i, l in enumerate(read_lines(path), 1):
                m = re.match(r'^rule_name = "(.*)"$', l)
                if m:
                    add("Battlefield rules", "Rule name (top right of the screen)", unescape(m.group(1)), path, i)


def scan_quips():
    path = "data/PythioLines.gd"
    moment = None
    for i, l in enumerate(read_lines(path), 1):
        if l.lstrip().startswith("#"):
            continue
        m = re.match(r'\s*"([A-Za-z_]+)":\s*\[', l)
        if m:
            moment = m.group(1)
            l = l[m.end():]
        for s in STRING.findall(l):
            add("Pythio's lines", "Pythio says this when " + MOMENTS.get(moment, moment or "?"), unescape(s), path, i)


def scan_code(path, area):
    is_scene = path.endswith(".tscn")
    for i, l in enumerate(read_lines(path), 1):
        stripped = l.strip()
        if stripped.startswith(("#", "@export_group", "@export_subgroup", "@export_category")) or not stripped:
            continue
        if is_scene:
            m = re.match(r'^(text|tooltip_text|placeholder_text) = "((?:[^"\\]|\\.)*)"$', l)
            if m and words_in(m.group(2)) and not re.fullmatch(r"[A-Za-z]+[A-Z][A-Za-z]*", m.group(2)):
                add(area, "On screen (%s; the game may replace it)" % ("tooltip" if m.group(1) == "tooltip_text" else "label"),
                    unescape(m.group(2)), path, i)
            continue
        code = re.sub(r'\s#\s.*$', "", l)          # drop trailing comments
        for s in STRING.findall(code):
            if is_player_text(s, code):
                where = "Tooltip" if "tooltip" in code else "On screen"
                add(area, where, unescape(s), path, i)


def main():
    try:
        import openpyxl
        from openpyxl.styles import Alignment, Font, PatternFill
        from openpyxl.worksheet.datavalidation import DataValidation  # noqa: F401 (checks the install)
    except ImportError:
        sys.exit("Needs openpyxl: python -m pip install openpyxl")

    scan_spells()
    scan_rules()
    scan_quips()
    for path, area in CODE_FILES.items():
        if os.path.exists(os.path.join(ROOT, path)):
            scan_code(path, area)
    for fn in sorted(os.listdir(os.path.join(ROOT, EFFECTS_DIR))):
        if fn.endswith(".gd"):
            scan_code(EFFECTS_DIR + "/" + fn, "Card rules text")

    rows.sort(key=lambda r: (AREAS.index(r["area"]), r["file"], r["line"]))

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Words"
    head = ["#", "Area", "Where it shows", "Current text", "Your version", "Notes", "File", "Line"]
    ws.append(head)
    for n, r in enumerate(rows, 1):
        ws.append([n, r["area"], r["where"], r["text"], "", "", r["file"], r["line"]])

    dark = PatternFill("solid", fgColor="2B2420")
    edit = PatternFill("solid", fgColor="FFF4D6")
    for c in ws[1]:
        c.font = Font(bold=True, color="F2E6D0")
        c.fill = dark
    widths = {"A": 6, "B": 20, "C": 42, "D": 60, "E": 60, "F": 30, "G": 34, "H": 7}
    for col, w in widths.items():
        ws.column_dimensions[col].width = w
    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.alignment = Alignment(wrap_text=True, vertical="top")
        row[4].fill = edit
        row[5].fill = edit
        for c in (row[6], row[7]):
            c.font = Font(color="8A8178", size=9)
    ws.freeze_panes = "E2"
    ws.auto_filter.ref = ws.dimensions

    how = wb.create_sheet("How to use", 0)
    guide = [
        "Pythio's words: every line a player reads",
        "",
        "1. Go to the Words tab. Filter the Area column to work through one part of the game at a time.",
        "2. Write your new text in the yellow 'Your version' column. Leave it empty to keep a line as it is.",
        "3. Use 'Notes' for anything else: an idea, 'cut this line', 'add another line like this one'.",
        "4. Save the file (keep it as .xlsx) and send it back. Claude applies every change in one PR.",
        "",
        "Things to keep:",
        "- Markers like %s, %d or {SHIELD} are filled in by the game (numbers, tag words). Keep them in your version.",
        "- [b], [color=...] and similar are formatting. Keep them, or tell Claude in Notes what you want.",
        "- Card rules text ('Heals 32.') is built from the numbers, so only its wording can change.",
        "- Renaming a spell renames it everywhere, including Pythio's hints and the leaderboard build line.",
        "",
        "Some lines may not be something you ever see (the export finds text automatically).",
        "Leave those empty, or write 'not shown' in Notes.",
        "Not included: the debug panel (F1). Schools are never shown, so their names aren't listed.",
        "",
        "Rebuild this file any time with: python tools/words/export_words.py",
    ]
    for line in guide:
        how.append([line])
    how["A1"].font = Font(bold=True, size=16)
    how.column_dimensions["A"].width = 110

    counts = wb.create_sheet("Counts")
    counts.append(["Area", "Lines"])
    counts["A1"].font = counts["B1"].font = Font(bold=True)
    for a in AREAS:
        counts.append([a, sum(1 for r in rows if r["area"] == a)])
    counts.append(["Total", len(rows)])
    counts.column_dimensions["A"].width = 26

    os.makedirs(os.path.dirname(OUT), exist_ok=True)
    wb.save(OUT)
    print("%d lines -> %s" % (len(rows), rel(OUT)))
    for a in AREAS:
        print("  %-22s %d" % (a, sum(1 for r in rows if r["area"] == a)))


if __name__ == "__main__":
    main()
